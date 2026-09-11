"""Tests de rendu pour les ajouts du brief v2 (questions/delta/epics_lecture) —
utilisent une Presentation python-pptx vide (aucune dépendance au template
réel du répertoire de données), voir conftest.py pour sys.path."""
from pptx import Presentation

import build_ppt
import render_agile


def _blank_slide_with_placeholder(prs, text="{{placeholder}}"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # layout "Blank"
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    box.text_frame.text = text
    return slide, box


def _texts(slide):
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame]


def _make_data(coaching):
    return {"project_name": "Projet Test", "coaching": coaching}


# ------------------------------------------------------- coach:decisions
def test_decisions_two_columns_when_delta_present():
    prs = Presentation()
    slide, ph = _blank_slide_with_placeholder(prs)
    data = _make_data({"delta": "Ce qui a bougé.", "questions": ["Une question ?"]})

    render_agile.render_coach_decisions(slide, ph, data, "unused.pptx", 1, 1)

    texts = " ".join(_texts(slide))
    assert "DEPUIS LE DERNIER POINT" in texts
    assert "QUESTIONS À OUVRIR" in texts
    assert "Ce qui a bougé." in texts
    assert "Une question ?" in texts


def test_decisions_one_column_when_delta_absent():
    prs = Presentation()
    slide, ph = _blank_slide_with_placeholder(prs)
    data = _make_data({"questions": ["Une question ?", "Une autre ?"]})

    render_agile.render_coach_decisions(slide, ph, data, "unused.pptx", 1, 1)

    texts = " ".join(_texts(slide))
    assert "DEPUIS LE DERNIER POINT" not in texts
    assert "QUESTIONS À OUVRIR" in texts
    assert "Une question ?" in texts
    assert "Une autre ?" in texts


def test_decisions_caps_at_three_questions():
    prs = Presentation()
    slide, ph = _blank_slide_with_placeholder(prs)
    questions = [f"Question numéro {i} ?" for i in range(5)]
    data = _make_data({"delta": "x", "questions": questions})

    render_agile.render_coach_decisions(slide, ph, data, "unused.pptx", 1, 1)

    texts = " ".join(_texts(slide))
    for q in questions[:3]:
        assert q in texts
    for q in questions[3:]:
        assert q not in texts


def test_decisions_fallback_message_when_both_absent():
    prs = Presentation()
    slide, ph = _blank_slide_with_placeholder(prs)
    data = _make_data({})

    render_agile.render_coach_decisions(slide, ph, data, "unused.pptx", 1, 1)

    texts = " ".join(_texts(slide))
    assert "Aucun delta ni question" in texts


# ---------------------------------------------------------- {{coach:decisions}} removal
def test_remove_slide_drops_it_from_presentation():
    prs = Presentation()
    slide_a = prs.slides.add_slide(prs.slide_layouts[6])
    slide_b = prs.slides.add_slide(prs.slide_layouts[6])
    assert len(prs.slides) == 2

    build_ppt._remove_slide(prs, slide_b)

    assert len(prs.slides) == 1
    assert prs.slides[0].slide_id == slide_a.slide_id


# --------------------------------------------------------------------- epics
def _make_epics_data(epics_lecture=None, n_items=10):
    items = [
        {
            "key": f"KEY-{i}",
            "name": f"Epic très très très long nom numéro {i} avec beaucoup de mots pour forcer le wrap",
            "status": "À faire",
            "estimate_jh": 10.0,
            "conso_jh": 5.0,
            "raf_theo_jh": 5.0,
            "us_done": 1,
            "us_remaining": 2,
            "avancement_pct": 33.3,
        }
        for i in range(n_items)
    ]
    epics = {
        "items": items,
        "us_done_total": sum(e["us_done"] for e in items),
        "total_estimate_jh": sum(e["estimate_jh"] for e in items),
        "total_conso_jh": sum(e["conso_jh"] for e in items),
        "total_raf_theo_jh": sum(e["raf_theo_jh"] for e in items),
        "avancement_pct": 33.3,
        "projection_theorique_jh": 40.0,
        "previsibilite_pct": 10.0,
    }
    coaching = {"epics_lecture": epics_lecture} if epics_lecture else {}
    return {
        "project_name": "Projet Test",
        "coaching": coaching,
        "metrics": {"epics": epics},
    }


def _table_shape(slide):
    for shape in slide.shapes:
        if shape.has_table:
            return shape
    return None


def test_epics_legend_rendered_when_present():
    prs = Presentation()
    slide, ph = _blank_slide_with_placeholder(prs)
    data = _make_epics_data(epics_lecture="Lecture des epics.")

    render_agile.render_epics(slide, ph, data, "unused.pptx", 1, 1)

    assert "Lecture des epics." in " ".join(_texts(slide))


def test_epics_table_shrinks_when_legend_present():
    prs1 = Presentation()
    slide1, ph1 = _blank_slide_with_placeholder(prs1)
    render_agile.render_epics(slide1, ph1, _make_epics_data(epics_lecture=None), "unused.pptx", 1, 1)
    table1 = _table_shape(slide1)

    prs2 = Presentation()
    slide2, ph2 = _blank_slide_with_placeholder(prs2)
    render_agile.render_epics(slide2, ph2, _make_epics_data(epics_lecture="Lecture."), "unused.pptx", 1, 1)
    table2 = _table_shape(slide2)

    assert table1 is not None and table2 is not None
    assert table2.height < table1.height
    assert "Lecture des epics" not in " ".join(_texts(slide1))


def test_epics_table_uses_original_height_when_legend_absent():
    prs = Presentation()
    slide, ph = _blank_slide_with_placeholder(prs)
    render_agile.render_epics(slide, ph, _make_epics_data(epics_lecture=None), "unused.pptx", 1, 1)

    table = _table_shape(slide)
    expected_h = render_agile.S.px(render_agile.S.CONTENT_H - 150 - 20)
    assert abs(table.height - expected_h) < render_agile.S.px(2)
