"""Helpers de construction de formes/texte pour les rendus de slide.

Toutes les coordonnées prises en paramètre sont en px de maquette (base
1920x1080) — la conversion en EMU/Pt passe systématiquement par style.py.
"""
import io
import zipfile

from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Pt

import style as S


def set_letter_spacing(run, spacing_pt):
    """Interlettrage (tracking) — pas d'API python-pptx dédiée, attribut XML spc
    (centièmes de pt), positif = espacé, négatif = resserré (approche)."""
    rPr = run._r.get_or_add_rPr()
    rPr.set("spc", str(round(spacing_pt * 100)))


def add_text(slide, x, y, w, h, text, size=S.SIZE_BODY, color=S.TEXT_PRIMARY,
             bold=False, align=PP_ALIGN.LEFT, font=S.FONT_FAMILY,
             anchor=MSO_ANCHOR.TOP, letter_spacing_pt=None, word_wrap=True,
             all_caps=False, min_size=S.MIN_CONTENT_PT):
    """Zone de texte simple, une seule ligne de run par paragraphe."""
    box = slide.shapes.add_textbox(S.px(x), S.px(y), S.px(w), S.px(h))
    tf = box.text_frame
    tf.word_wrap = word_wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text.upper() if all_caps else text
    run.font.size = size if size.pt >= min_size.pt else min_size
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color
    if letter_spacing_pt is not None:
        set_letter_spacing(run, letter_spacing_pt)
    return box


def add_rounded_rect(slide, x, y, w, h, fill=S.CARD_BG, line_color=S.CARD_BORDER,
                      line_width_pt=1, radius_px=S.CARD_RADIUS_PX, shadow=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, S.px(x), S.px(y), S.px(w), S.px(h))
    try:
        adj = radius_px / min(w, h) if min(w, h) else 0.08
        shape.adjustments[0] = min(max(adj, 0.0), 0.5)
    except Exception:
        pass
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_width_pt)
    shape.shadow.inherit = False
    return shape


def add_rect(slide, x, y, w, h, fill=S.CARD_BG, line_color=None, line_width_pt=1):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, S.px(x), S.px(y), S.px(w), S.px(h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_width_pt)
    shape.shadow.inherit = False
    return shape


def set_transparency(fill_or_shape, pct):
    """Opacité de remplissage (0-100 = % transparent) — pas d'API python-pptx,
    injection de l'élément a:alpha dans le srgbClr du fill. Accepte soit un
    FillFormat (shape.fill / series.format.fill) soit directement une shape."""
    fill = fill_or_shape.fill if hasattr(fill_or_shape, "fill") else fill_or_shape
    sp = fill.fore_color._xFill
    clr = sp.find(qn("a:srgbClr"))
    if clr is None:
        return
    alpha = clr.makeelement(qn("a:alpha"), {"val": str(round((100 - pct) * 1000))})
    clr.append(alpha)


_TEMPLATE_IMAGE_CACHE = {}


def template_image_bytes(template_path, media_name):
    """Extrait un média (ex. 'image3.png') du .pptx source pour le
    réutiliser (logo, fond de couverture) sans dupliquer les binaires."""
    key = (str(template_path), media_name)
    if key not in _TEMPLATE_IMAGE_CACHE:
        with zipfile.ZipFile(template_path) as z:
            _TEMPLATE_IMAGE_CACHE[key] = z.read(f"ppt/media/{media_name}")
    return io.BytesIO(_TEMPLATE_IMAGE_CACHE[key])


def add_header(slide, surtitre, titre, template_path, logo_media="image3.png"):
    add_text(slide, S.MARGIN_X, S.MARGIN_TOP, S.CONTENT_W - S.LOGO_WIDTH_PX - 20, 26,
              surtitre, size=S.SIZE_SURTITLE, color=S.BRAND_ACCENT, bold=True,
              all_caps=True, letter_spacing_pt=1.2)
    add_text(slide, S.MARGIN_X, S.MARGIN_TOP + 33, S.CONTENT_W - S.LOGO_WIDTH_PX - 20, 62,
              titre, size=S.SIZE_SLIDE_TITLE, color=S.BRAND_PRIMARY, bold=True,
              letter_spacing_pt=-0.3)
    logo_h = S.LOGO_WIDTH_PX * 0.263  # ratio approx image3.png (230x60ish)
    slide.shapes.add_picture(
        template_image_bytes(template_path, logo_media),
        S.px(S.MARGIN_X + S.CONTENT_W - S.LOGO_WIDTH_PX), S.px(S.MARGIN_TOP),
        S.px(S.LOGO_WIDTH_PX),
    )


def add_footer(slide, project_name, page_num, total_pages):
    y = S.SLIDE_H_PX - S.MARGIN_BOTTOM + 22
    add_text(slide, S.MARGIN_X, y, 600, 30, project_name,
              size=S.SIZE_FOOTER, color=S.TEXT_FOOTER)
    add_text(slide, S.MARGIN_X + S.CONTENT_W - 150, y, 150, 30,
              f"{page_num:02d} / {total_pages}", size=S.SIZE_FOOTER,
              color=S.TEXT_FOOTER, align=PP_ALIGN.RIGHT)


def clear_placeholder(shape):
    """Retire du slide la zone de texte {{...}} d'origine une fois son
    contenu extrait (nom de slide, position...)."""
    shape._element.getparent().remove(shape._element)


def clear_legacy_chrome(slide):
    """Un rendu « design agile » dessine son propre header/footer (surtitre,
    titre, logo, pied de page) — retire les éléments du template simple qui
    feraient doublon : la zone de texte {{project_name}}, le placeholder
    natif de numéro de slide, et le logo déjà présent en haut à droite."""
    for shape in list(slide.shapes):
        is_project_name = shape.has_text_frame and "{{project_name}}" in shape.text_frame.text
        is_slide_number = (
            shape.is_placeholder
            and shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
        )
        is_header_logo = (
            shape.shape_type == 13  # PICTURE
            and shape.top is not None
            and shape.top < S.px(300)
        )
        if is_project_name or is_slide_number or is_header_logo:
            shape._element.getparent().remove(shape._element)
