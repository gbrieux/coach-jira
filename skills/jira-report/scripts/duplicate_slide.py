#!/usr/bin/env python3
"""Duplique une slide de template.pptx pour y placer un nouveau placeholder
{{chart:xxx}} — utilitaire réutilisable à chaque ajout d'indicateur (voir
SKILL.md, section « Ajouter un indicateur »).

Chaque slide du template suit le même moule : {{project_name}}, un placeholder
de contenu ({{chart:...}}, {{liste_...}}, {{kpi:*}}, {{coach:synthese}}), un
numéro de page en texte brut (pas un champ auto-calculé). Dupliquer une slide
à la main dans PowerPoint est fastidieux et casse facilement la numérotation
des pages suivantes — ce script fait les trois choses ensemble :

1. Trouve la slide source (celle qui contient `--source-placeholder`) et la
   duplique intégralement (mise en forme comprise), copie insérée juste après
   l'original.
2. Remplace le texte du placeholder sur la copie par `--new-placeholder`.
3. Renumérote tous les numéros de page en texte brut du deck (1..N) pour
   rester cohérent après l'insertion.

Usage:
    python scripts/duplicate_slide.py <template.pptx> \\
        --source-placeholder "{{chart:types_conso}}" \\
        --new-placeholder "{{chart:conso_corrective}}"

Modifie le fichier en place (pas de sauvegarde séparée — le template est déjà
versionné dans git, voir SKILL.md).
"""
import argparse
import copy
import re
import sys

from pptx import Presentation

PAGE_NUM_RE = re.compile(r"^\d+$")


def die(msg):
    print(f"ERREUR: {msg}", file=sys.stderr)
    sys.exit(1)


def find_slide_with_text(prs, needle):
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame and needle in shape.text_frame.text:
                return i
    return None


def duplicate_slide(prs, source_index):
    """Copie la slide `source_index` (0-based) à la fin du deck, mise en
    forme comprise (mise en page, position, style de chaque forme), puis la
    remonte juste après l'original dans l'ordre. Retourne la nouvelle slide."""
    source = prs.slides[source_index]
    dest = prs.slides.add_slide(source.slide_layout)
    for shp in list(dest.shapes):  # vide les placeholders auto-ajoutés par la layout
        shp._element.getparent().remove(shp._element)
    for shp in source.shapes:
        dest.shapes._spTree.append(copy.deepcopy(shp._element))

    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    new_elem = slides[-1]
    xml_slides.remove(new_elem)
    xml_slides.insert(source_index + 1, new_elem)
    return prs.slides[source_index + 1]


def replace_shape_text(slide, old_text, new_text):
    """Remplace `old_text` par `new_text` dans la première forme qui le
    contient, run par run (comme build_ppt.py:replace_runs) pour préserver la
    mise en forme existante."""
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        tf = shape.text_frame
        if old_text not in tf.text:
            continue
        for para in tf.paragraphs:
            for run in para.runs:
                if old_text in run.text:
                    run.text = run.text.replace(old_text, new_text)
            full = "".join(r.text for r in para.runs)
            if old_text in full and para.runs:
                para.runs[0].text = full.replace(old_text, new_text)
                for r in para.runs[1:]:
                    r.text = ""
        return True
    return False


def renumber_pages(prs):
    """Remet à plat les numéros de page en texte brut (1..N) sur tout le
    deck, dans l'ordre des slides — un numéro par slide, seule forme dont le
    texte est purement numérique."""
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame and PAGE_NUM_RE.match(shape.text_frame.text.strip()):
                tf = shape.text_frame
                para = tf.paragraphs[0]
                if para.runs:
                    para.runs[0].text = str(i)
                    for r in para.runs[1:]:
                        r.text = ""
                else:
                    tf.text = str(i)
                break


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("template", help="chemin du template.pptx à modifier (en place)")
    ap.add_argument("--source-placeholder", required=True,
                     help="placeholder de la slide à dupliquer, ex. {{chart:types_conso}}")
    ap.add_argument("--new-placeholder", required=True,
                     help="placeholder à écrire sur la copie, ex. {{chart:conso_corrective}}")
    args = ap.parse_args()

    prs = Presentation(args.template)
    idx = find_slide_with_text(prs, args.source_placeholder)
    if idx is None:
        die(f"Aucune slide ne contient '{args.source_placeholder}' dans {args.template}.")

    new_slide = duplicate_slide(prs, idx)
    if not replace_shape_text(new_slide, args.source_placeholder, args.new_placeholder):
        die("Slide dupliquée mais placeholder introuvable dessus (bug du script).")
    renumber_pages(prs)

    prs.save(args.template)
    print(f"OK: slide {idx + 2} insérée dans {args.template} avec {args.new_placeholder} "
          f"(dupliquée depuis la slide {idx + 1}). Vérifier avec inspect_template.py.")


if __name__ == "__main__":
    main()
