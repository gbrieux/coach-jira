#!/usr/bin/env python3
"""Génère le PPT d'indicateurs à partir du JSON JIRA et du template configuré.

Usage:
    python scripts/build_ppt.py <clé>
    python scripts/build_ppt.py --all

Placeholders texte {{...}} remplacés run par run (formatage du template préservé).
Placeholders graphiques {{chart:xxx}} remplacés par des graphiques natifs.
"""
import argparse
import json
import re
import sys
from datetime import datetime

import indicators
import render_agile
from pptx import Presentation

from workspace import output_dir, template_path, list_project_keys, DataDirNotConfigured

PLACEHOLDER = re.compile(r"\{\{[^}]+\}\}")
CHART_PH = re.compile(r"\{\{chart:(\w+)\}\}")
TABLE_PH = "{{liste_sprints}}"
EPICS_PH = "{{liste_epics}}"
COACH_PH = "{{coach:synthese}}"
KPI_PH = "{{kpi:*}}"

# {placeholder: render(slide, shape, data, template_path, page_num, total_pages)}
# — découvert automatiquement depuis indicators/ (voir indicators/__init__.py).
# Ajouter un indicateur n'implique aucune modification ici.
AGILE_CHART_RENDERERS = {}
for _module in indicators.all_modules():
    AGILE_CHART_RENDERERS.update(getattr(_module, "RENDERERS", {}))

SPRINT_TABLE_HEADERS = [
    "Sprint", "Num sprint", "Date début", "Date fin",
    "Nb de US terminées", "US terminées cumulées",
    "Nb de US ajoutées", "Périmètre cumulé US (nb)",
    "Estimation US", "Périmètre cumulé US (jh)",
    "Conso US terminées", "Conso cumulée US (jh)",
]


def die(msg):
    print(f"ERREUR: {msg}", file=sys.stderr)
    sys.exit(1)


def load_data(key):
    p = output_dir(key) / "data.json"
    if not p.exists():
        die(f"{p} manquant. Lancer d'abord: python scripts/fetch_jira.py {key}")
    return json.loads(p.read_text(encoding="utf-8"))


def _iter_text_values(value):
    """Parcourt récursivement les chaînes d'une structure JSON."""
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for key, item in value.items():
            yield from _iter_text_values(key)
            yield from _iter_text_values(item)
    elif isinstance(value, list):
        for item in value:
            yield from _iter_text_values(item)


def validate_coaching_encoding(data, key):
    """Refuse un coaching probablement corrompu par un pipe ASCII PowerShell.

    Windows PowerShell 5.1 utilise ``$OutputEncoding=us-ascii`` par défaut pour
    les programmes natifs. Un bloc français envoyé avec ``... | python -``
    transforme donc les caractères accentués en ``?`` avant même que Python ne
    reçoive le texte. Le PPT ne peut pas reconstruire ces caractères perdus.
    """
    coaching = data.get("coaching")
    if not coaching:
        return

    texts = list(_iter_text_values(coaching))
    replacement_chars = sum(text.count("\ufffd") for text in texts)
    question_marks = sum(text.count("?") for text in texts)
    embedded_question_mark = any(
        re.search(r"[^\W\d_]\?[^\W\d_]", text, re.UNICODE) for text in texts
    )

    if replacement_chars or embedded_question_mark or question_marks > 3:
        die(
            f"[{key}] texte de coaching probablement corrompu par l'encodage "
            f"({question_marks} caractère(s) '?' et {replacement_chars} caractère(s) de remplacement). "
            "Réinjecter le coaching avec scripts/set_coaching.py depuis un fichier JSON UTF-8. "
            "Sous Windows PowerShell, ne pas utiliser un here-string accentué avec '| python -'."
        )


def short_sprint(name):
    """'Nouvelle Offre Santé Sprint 12' -> 'S12' pour des axes lisibles."""
    m = re.search(r"(\d+)\s*$", name or "")
    return f"S{m.group(1)}" if m else (name or "")[:10]


def format_date(iso):
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(str(iso).replace("Z", "+00:00")).strftime("%d/%m/%Y")
    except ValueError:
        return ""


def sprint_table_rows(data):
    """Une ligne par sprint (dont le Sprint 0 = backlog) ayant eu de
    l'activité côté User Story (ajout ou clôture)."""
    rows = []
    for b in data["metrics"].get("sprint_table", []):
        if not (b.get("added_us") or b.get("done_us")):
            continue
        rows.append([
            short_sprint(b["sprint"]),
            b.get("num") if b.get("num") is not None else "",
            format_date(b.get("start")),
            format_date(b.get("end")),
            b.get("done_us", 0),
            b.get("done_cumul_us", 0),
            b.get("added_us", 0),
            b.get("scope_cumul_us", 0),
            b.get("estimate_us_jh", 0),
            b.get("scope_cumul_us_jh", 0),
            b.get("conso_us_jh", 0),
            b.get("conso_cumul_us_jh", 0),
        ])
    return rows


def text_map(data):
    m = data["metrics"]
    ct = m["cycle_time"]
    return {
        "{{project_name}}": data["project_name"],
        "{{project_key}}": data["project_key"],
        "{{date}}": data["generated_at"][:10],
        "{{jql}}": data["jql"],
        "{{total_issues}}": str(m["total_issues"]),
        "{{total_story_points}}": str(m["total_story_points"]),
        "{{avg_velocity}}": str(m["avg_velocity"]),
        "{{nb_sprints}}": str(len(m["sprints"])),
        "{{kpi:cycle_median}}": f"{ct['median']} j",
        "{{kpi:cycle_mean}}": f"{ct['mean']} j",
        "{{kpi:cycle_p15}}": f"{ct['p15']} j",
        "{{kpi:cycle_p85}}": f"{ct['p85']} j",
        "{{kpi:cycle_count}}": str(ct["count"]),
        "{{kpi:*}}": (
            f"Cycle time médian : {ct['median']} j  |  "
            f"Moyenne : {ct['mean']} j  |  "
            f"P15-P85 : {ct['p15']}-{ct['p85']} j  |  "
            f"US mesurées : {ct['count']}"
        ),
    }


def replace_runs(tf, mapping):
    for para in tf.paragraphs:
        for run in para.runs:
            for k, v in mapping.items():
                if k in run.text:
                    run.text = run.text.replace(k, v)
        full = "".join(r.text for r in para.runs)
        if PLACEHOLDER.search(full) and para.runs:
            new = full
            for k, v in mapping.items():
                new = new.replace(k, v)
            if PLACEHOLDER.search(new):
                new = PLACEHOLDER.sub("", new)
            para.runs[0].text = new
            for r in para.runs[1:]:
                r.text = ""


def process(key):
    try:
        TEMPLATE = template_path()
    except DataDirNotConfigured as e:
        die(str(e))
    if not TEMPLATE.exists():
        die(f"Template absent: {TEMPLATE}. Le déposer dans templates/ (voir SKILL.md).")
    data = load_data(key)
    validate_coaching_encoding(data, key)
    prs = Presentation(str(TEMPLATE))
    mapping = text_map(data)

    # Couverture (slide 1) : pas de placeholder {{...}}, positionnelle.
    if len(prs.slides) > 0:
        render_agile.render_cover(prs.slides[0], data)

    placed = []
    # Snapshot avant boucle : render_coach peut ajouter une slide (débordement
    # >2000 caractères) — itérer sur une liste figée évite de la revisiter.
    for slide in list(prs.slides):
        chart_shapes, table_shapes, epics_shapes, coach_shapes, kpi_shapes = [], [], [], [], []
        for shape in slide.shapes:
            if shape.has_text_frame:
                mt = CHART_PH.search(shape.text_frame.text)
                if mt:
                    chart_shapes.append((shape, mt.group(1)))
                elif TABLE_PH in shape.text_frame.text:
                    table_shapes.append(shape)
                elif EPICS_PH in shape.text_frame.text:
                    epics_shapes.append(shape)
                elif COACH_PH in shape.text_frame.text:
                    coach_shapes.append(shape)
                elif KPI_PH in shape.text_frame.text:
                    kpi_shapes.append(shape)

        for shape, cname in chart_shapes:
            renderer = AGILE_CHART_RENDERERS.get(cname)
            if renderer:
                page_num = list(prs.slides).index(slide) + 1
                renderer(slide, shape, data, TEMPLATE, page_num, len(prs.slides))
                placed.append(cname)
            else:
                print(f"[{key}] AVERTISSEMENT: chart '{cname}' non reconnu (aucun rendu associé)")

        for shape in table_shapes:
            page_num = list(prs.slides).index(slide) + 1
            render_agile.render_sprint_table(slide, shape, data, TEMPLATE, page_num, len(prs.slides))
            placed.append("liste_sprints")

        for shape in epics_shapes:
            page_num = list(prs.slides).index(slide) + 1
            render_agile.render_epics(slide, shape, data, TEMPLATE, page_num, len(prs.slides))
            placed.append("liste_epics")

        for shape in kpi_shapes:
            page_num = list(prs.slides).index(slide) + 1
            render_agile.render_kpi(slide, shape, data, TEMPLATE, page_num, len(prs.slides))
            placed.append("kpi")

        for shape in coach_shapes:
            page_num = list(prs.slides).index(slide) + 1
            render_agile.render_coach(slide, shape, data, TEMPLATE, page_num, len(prs.slides), prs=prs)
            if not data.get("coaching"):
                print(f"[{key}] AVERTISSEMENT: coach:synthese demandé, pas de synthèse dans le JSON "
                      f"(clé 'coaching' absente — voir COACH_PROMPT.md)")
            placed.append("coach:synthese")

        for shape in slide.shapes:
            if shape.has_text_frame:
                replace_runs(shape.text_frame, mapping)

    out = output_dir(key) / "indicateurs.pptx"
    prs.save(str(out))
    print(f"[{key}] écrit {out}  (slides: {', '.join(placed) or 'aucune'})")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("project_key", nargs="?")
    ap.add_argument("--all", action="store_true")
    args = ap.parse_args()
    if args.all:
        keys = list_project_keys()
        if not keys:
            die("Aucun projet dans projects/. Lancer fetch_jira.py --all d'abord.")
        for k in keys:
            process(k)
    elif args.project_key:
        process(args.project_key)
    else:
        die("Préciser une clé projet ou --all.")


if __name__ == "__main__":
    main()
