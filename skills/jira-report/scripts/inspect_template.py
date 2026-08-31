#!/usr/bin/env python3
"""Liste les placeholders {{...}} et les zones de graphiques d'un template.

Usage: python scripts/inspect_template.py [chemin_template.pptx]
Sans argument, inspecte templates/template.pptx dans le répertoire de données.
"""
import re
import sys
from pathlib import Path

from pptx import Presentation

from workspace import template_path, DataDirNotConfigured

PLACEHOLDER = re.compile(r"\{\{[^}]+\}\}")


def main():
    if len(sys.argv) > 1:
        tpl = Path(sys.argv[1])
    else:
        try:
            tpl = template_path()
        except DataDirNotConfigured as e:
            print(f"ERREUR: {e}", file=sys.stderr)
            sys.exit(1)
    if not tpl.exists():
        print(f"ERREUR: template introuvable: {tpl}", file=sys.stderr)
        sys.exit(1)

    prs = Presentation(str(tpl))
    found = set()
    print(f"Template: {tpl}\n{'='*60}")
    for i, slide in enumerate(prs.slides, 1):
        print(f"\n--- Slide {i} ---")
        for shape in slide.shapes:
            kind = shape.shape_type
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                phs = PLACEHOLDER.findall(txt)
                for p in phs:
                    found.add(p)
                label = txt[:50].replace("\n", " ")
                if phs:
                    print(f"  [texte] '{label}'  -> {', '.join(phs)}")
                elif label:
                    print(f"  [texte] '{label}'")
            elif shape.has_chart:
                print(f"  [chart existant] name='{shape.name}'")
            else:
                print(f"  [{kind}] name='{shape.name}'")

    print(f"\n{'='*60}\nPlaceholders détectés ({len(found)}):")
    for p in sorted(found):
        print(f"  {p}")
    if not found:
        print("  AUCUN. Ajouter des {{...}} dans les zones de texte du template")
        print("  (voir la table des indicateurs dans SKILL.md).")


if __name__ == "__main__":
    main()
